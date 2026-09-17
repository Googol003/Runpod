"""
Gezielte, sparsame Test-Injects in test.1175.xlsx.

Fehlerquellen:
1) TEXT_LEAK x3 — Wort vom vorherigen Sprecher klebt am nächsten Segment
2) NONSENSE x1 — nahe Segmente, Kontext-Unsinn
3) SPEAKER_WRONG x2 — MATCHED-TEXT/SOURCE/REF aus Nachbarzeile
4) MUSCHI-Split — nur „Du unglaubliche.“, „Muschi“ weg
5) OVERLAP-CHAOS-Block (32–36) — Fokus falsche Rolle + falscher Text;
   eine Zeile unmatched → Original in NOT_MATCHED

Fehlerzeilen: Spalte INJECTED-ERROR + rote Zellenfüllung.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

KEEP = [
    "TIMECODE-IN",
    "TIMECODE-OUT",
    "DIALOGUE",
    "SOURCE",
    "MATCHED-TEXT",
    "REF-IN",
    "REF-OUT",
    "NOT_MATCHED",
]

RED = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")


def _c(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("nan", "none") else s


def _copy_match(out: pd.DataFrame, dst: int, src: int) -> None:
    out.at[dst, "SOURCE"] = out.at[src, "SOURCE"]
    out.at[dst, "MATCHED-TEXT"] = out.at[src, "MATCHED-TEXT"]
    out.at[dst, "REF-IN"] = out.at[src, "REF-IN"]
    out.at[dst, "REF-OUT"] = out.at[src, "REF-OUT"]


def _set_match(
    out: pd.DataFrame,
    dst: int,
    *,
    source: str,
    matched: str,
    ref_in: str,
    ref_out: str,
) -> None:
    out.at[dst, "SOURCE"] = source
    out.at[dst, "MATCHED-TEXT"] = matched
    out.at[dst, "REF-IN"] = ref_in
    out.at[dst, "REF-OUT"] = ref_out


def _nm_line(*, script_i: int, tc_in: str, tc_out: str, speaker: str, text: str) -> str:
    return (
        f"[Original ohne Transkript-Match] (Script-Zeile #{script_i}) "
        f"Original: {tc_in} – {tc_out} | {speaker}: {text} | "
        f"Platzierung: nicht prüfbar (kein Refine-Zeitcode, Drehbuch-Matcher fehlt oder keine Zeile)"
    )


def inject(df: pd.DataFrame) -> tuple[pd.DataFrame, dict[int, str]]:
    out = df.copy()
    notes: dict[int, str] = {}

    # --- TEXT_LEAK x3 ---
    out.at[4, "DIALOGUE"] = "hast – Ich hab dich nie gezwungen, ihn zu verlassen."
    notes[4] = (
        "TEXT_LEAK: Wort „hast“ vom vorherigen STEDE-Segment klebt am IZZY-Satz "
        "(mit anderem Satzzeichen „–“)."
    )

    out.at[31, "DIALOGUE"] = "Mann! Er hat mich angefurzt."
    notes[31] = (
        "TEXT_LEAK: Wort „Mann“ vom vorherigen BUTTONS-Segment klebt am STEDE-Satz "
        "(Satzzeichen „!“ statt „?“)."
    )

    out.at[56, "DIALOGUE"] = "du? Das ist der Schwede."
    notes[56] = (
        "TEXT_LEAK: Wort „du“ vom vorherigen STEDE-Segment klebt am OLUWANDE-Satz "
        "(„?“); zusätzlich SPEAKER_WRONG (siehe Match)."
    )

    # --- NONSENSE x1 (sehr nahe an „Wo ist er?“) ---
    out.at[7, "DIALOGUE"] = "Wo ist das Sofa?"
    notes[7] = (
        "NONSENSE: Sehr nahe am vorherigen Segment „Wo ist er?“ — "
        "„Wo ist das Sofa?“ ergibt im Kontext (Suche nach Ed) keinen Sinn."
    )

    # --- SPEAKER_WRONG / falscher MATCHED-TEXT (einzeln) ---
    _copy_match(out, dst=10, src=9)
    notes[10] = (
        "SPEAKER_WRONG: Dialog „Ed!“ behält, aber MATCHED-TEXT/SOURCE/REF "
        "fälschlich von vorherigem BLACKBEARD-„Stede!“-Segment übernommen."
    )

    _copy_match(out, dst=56, src=57)
    notes[56] = (
        notes[56]
        + " SPEAKER_WRONG: MATCHED-TEXT/SOURCE/REF fälschlich vom nächsten "
        "SCHWEDE-Segment („Ich bin der Schwede.“) übernommen."
    )

    # --- „Du unglaubliche“ nur einmal ---
    out.at[8, "DIALOGUE"] = "Du unglaubliche."
    notes[8] = (
        "ASR/SPLIT: Nur „Du unglaubliche.“ belassen — „Muschi“ entfernt "
        "(kein doppelter Anfang). Match bleibt Kackvogel-Zeile."
    )

    # --- OVERLAP-CHAOS 32–36: Fokus falsche Rolle + falscher Text ---
    real_32 = {
        "source": _c(out.at[32, "SOURCE"]),
        "matched": _c(out.at[32, "MATCHED-TEXT"]),
        "ref_in": _c(out.at[32, "REF-IN"]),
        "ref_out": _c(out.at[32, "REF-OUT"]),
    }
    real_34 = {
        "matched": _c(out.at[34, "MATCHED-TEXT"]),
        "ref_in": _c(out.at[34, "REF-IN"]),
        "ref_out": _c(out.at[34, "REF-OUT"]),
    }
    real_36 = {
        "matched": _c(out.at[36, "MATCHED-TEXT"]),
        "ref_in": _c(out.at[36, "REF-IN"]),
        "ref_out": _c(out.at[36, "REF-OUT"]),
    }
    buttons_30 = {
        "source": _c(out.at[30, "SOURCE"]),
        "matched": _c(out.at[30, "MATCHED-TEXT"]),
        "ref_in": _c(out.at[30, "REF-IN"]),
        "ref_out": _c(out.at[30, "REF-OUT"]),
    }

    # 32: ROACH+Pete vermischt → fälschlich STEDE + Gestank-Match
    out.at[32, "DIALOGUE"] = "Halt die Klappe Ed oh Ed die Nacht!"
    _set_match(
        out,
        32,
        source="STEDE",
        matched=real_34["matched"],
        ref_in=real_34["ref_in"],
        ref_out=real_34["ref_out"],
    )
    notes[32] = (
        "OVERLAP-CHAOS / SPEAKER_WRONG+TEXT: Dialog mischt ROACH+PETE; "
        "fälschlich SOURCE=STEDE und MATCHED-TEXT=Stede-Gestank-Zeile."
    )

    # 33: Pete/Stede-Wörter → fälschlich BUTTONS
    out.at[33, "DIALOGUE"] = "Schnauze Wand Gestank"
    _set_match(
        out,
        33,
        source=buttons_30["source"],
        matched=buttons_30["matched"],
        ref_in=buttons_30["ref_in"],
        ref_out=buttons_30["ref_out"],
    )
    notes[33] = (
        "OVERLAP-CHAOS / SPEAKER_WRONG+TEXT: „Schnauze Wand Gestank“; "
        "fälschlich SOURCE=BUTTONS + Buttons-MATCHED-TEXT."
    )

    # 34: unmatched → Original in NOT_MATCHED
    existing_nm = _c(out.at[34, "NOT_MATCHED"])
    stede_nm = _nm_line(
        script_i=29,
        tc_in=real_34["ref_in"] or "00:02:28:02",
        tc_out="00:02:31:02",
        speaker="STEDE",
        text=real_34["matched"] or "Es ist noch da. Es ist 'ne Wand aus Gestank.",
    )
    out.at[34, "DIALOGUE"] = "Es ist noch da irgendwie die Klappe"
    out.at[34, "SOURCE"] = ""
    out.at[34, "MATCHED-TEXT"] = ""
    out.at[34, "REF-IN"] = ""
    out.at[34, "REF-OUT"] = ""
    out.at[34, "NOT_MATCHED"] = (
        f"{existing_nm} || {stede_nm}" if existing_nm else stede_nm
    )
    notes[34] = (
        "OVERLAP-CHAOS / UNMATCHED: STEDE ohne Match (SOURCE/MATCHED/REF leer); "
        "echtes Original in NOT_MATCHED — am Original orientieren/korrigieren."
    )

    # 35: wirres ASR → fälschlich PETE-Match
    out.at[35, "DIALOGUE"] = "Es ist eine Wand aus Ed oh Ed Gestöhne"
    _set_match(
        out,
        35,
        source="BLACK PETE",
        matched=real_36["matched"],
        ref_in=real_36["ref_in"],
        ref_out=real_36["ref_out"],
    )
    notes[35] = (
        "OVERLAP-CHAOS / SPEAKER_WRONG+TEXT: Gestank+Pete vermischt; "
        "fälschlich SOURCE=BLACK PETE + Pete-MATCHED-TEXT."
    )

    # 36: Pete+Roach → fälschlich ROACH-Match
    out.at[36, "DIALOGUE"] = "Kann nicht schlimmer Halt die Klappe die ganze Nacht"
    _set_match(
        out,
        36,
        source=real_32["source"] or "ROACH",
        matched=real_32["matched"] or "Halt die Klappe!",
        ref_in=real_32["ref_in"],
        ref_out=real_32["ref_out"],
    )
    notes[36] = (
        "OVERLAP-CHAOS / SPEAKER_WRONG+TEXT: PETE-Satz mit ROACH-Leak; "
        "fälschlich SOURCE=ROACH + MATCHED-TEXT „Halt die Klappe!“."
    )

    extra_rows = [
        {
            "TIMECODE-IN": "00:04:46:02",
            "TIMECODE-OUT": "00:04:48:10",
            "DIALOGUE": "Darf ich vorstellen: Bjungho Kim und Sungmin Bark.",
            "SOURCE": "STEDE",
            "MATCHED-TEXT": "Darf ich vorstellen: Byung-ho Kim und Seung-min Park.",
            "REF-IN": "00:04:46:02",
            "REF-OUT": "00:04:48:12",
            "NOT_MATCHED": "",
            "INJECTED-ERROR": (
                "NAME_ERROR/PHONETIC: Byung-ho Kim → Bjungho Kim; "
                "Seung-min Park → Sungmin Bark. Original-Schreibweise nutzen."
            ),
        },
        {
            "TIMECODE-IN": "00:04:48:14",
            "TIMECODE-OUT": "00:04:49:20",
            "DIALOGUE": "Und das ist Kinghi Tschoi.",
            "SOURCE": "STEDE",
            "MATCHED-TEXT": "Und das ist Gyeong-hui Choi.",
            "REF-IN": "00:04:48:14",
            "REF-OUT": "00:04:49:22",
            "NOT_MATCHED": "",
            "INJECTED-ERROR": (
                "NAME_ERROR/PHONETIC: Gyeong-hui Choi → Kinghi Tschoi."
            ),
        },
        {
            "TIMECODE-IN": "00:04:50:00",
            "TIMECODE-OUT": "00:04:51:18",
            "DIALOGUE": "Pingho hier, Captain Bonnet.",
            "SOURCE": "BYUNG-HO",
            "MATCHED-TEXT": "Freut mich, Captain Bonnet.",
            "REF-IN": "00:04:50:00",
            "REF-OUT": "00:04:51:20",
            "NOT_MATCHED": "",
            "INJECTED-ERROR": (
                "NAME_ERROR + NONSENSE: Byung-ho komplett falsch als Pingho; "
                "Satzanfang nicht Original. Orig: Freut mich, Captain Bonnet."
            ),
        },
        {
            "TIMECODE-IN": "00:04:52:00",
            "TIMECODE-OUT": "00:04:54:08",
            "DIALOGUE": "Gion wartet draußen bei Cheng Isao.",
            "SOURCE": "SEUNG-MIN",
            "MATCHED-TEXT": "Ji-yeon wartet draußen bei Zheng Yi Sao.",
            "REF-IN": "00:04:52:00",
            "REF-OUT": "00:04:54:10",
            "NOT_MATCHED": "",
            "INJECTED-ERROR": (
                "NAME_ERROR: Ji-yeon → Gion; Zheng Yi Sao → Cheng Isao."
            ),
        },
        {
            "TIMECODE-IN": "00:04:54:12",
            "TIMECODE-OUT": "00:04:55:22",
            "DIALOGUE": "Uschin kommt gleich, Captain.",
            "SOURCE": "GYEONG-HUI",
            "MATCHED-TEXT": "Woo-jin kommt gleich, Captain.",
            "REF-IN": "00:04:54:12",
            "REF-OUT": "00:04:56:00",
            "NOT_MATCHED": "",
            "INJECTED-ERROR": "NAME_ERROR/PHONETIC: Woo-jin → Uschin.",
        },
        {
            "TIMECODE-IN": "00:04:56:02",
            "TIMECODE-OUT": "00:04:57:16",
            "DIALOGUE": "Joey Jin, nicht wahr?",
            "SOURCE": "STEDE",
            "MATCHED-TEXT": "Woo-jin Jeong, nicht wahr?",
            "REF-IN": "00:04:56:02",
            "REF-OUT": "00:04:57:18",
            "NOT_MATCHED": "",
            "INJECTED-ERROR": (
                "NAME_ERROR komplett falsch: Woo-jin Jeong → Joey Jin "
                "(ASR hat den Namen erfunden)."
            ),
        },
    ]
    extra = pd.DataFrame(extra_rows)
    start_i = len(out)
    out = pd.concat([out, extra], ignore_index=True)
    for j, row in enumerate(extra_rows):
        notes[start_i + j] = row["INJECTED-ERROR"]

    out["INJECTED-ERROR"] = ""
    for i, note in notes.items():
        out.at[i, "INJECTED-ERROR"] = note

    return out, notes


def _paint_red(path: Path, error_rows: set[int]) -> None:
    wb = load_workbook(path)
    ws = wb.active
    for df_i in error_rows:
        excel_row = df_i + 2
        for col in range(1, ws.max_column + 1):
            ws.cell(excel_row, col).fill = RED
    wb.save(path)


def main() -> None:
    base = Path(__file__).resolve().parent
    clean = base / "output" / "test.1175.relevant_100.xlsx"
    if not clean.exists():
        raise SystemExit(f"Clean base missing: {clean}")
    df = pd.read_excel(clean)
    df = df[[c for c in KEEP if c in df.columns]].iloc[:100].copy()

    injected, notes = inject(df)
    error_rows = set(notes.keys())

    outs = (
        base / "testdata" / "test.1175.xlsx",
        base / "testdata" / "test.1175.injected.xlsx",
        base / "output" / "test.1175.injected.xlsx",
    )
    for out in outs:
        out.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(out, engine="openpyxl") as w:
            injected.to_excel(w, sheet_name="Transkription", index=False)
        _paint_red(out, error_rows)
        print(f"wrote {out} rows={len(injected)} red_rows={sorted(error_rows)}")

    print("--- OVERLAP 32-36 ---")
    for i in range(32, 37):
        r = injected.iloc[i]
        print(
            f"  [{i}] SRC={r['SOURCE']!r} D={str(r['DIALOGUE'])[:50]!r} "
            f"M={str(r['MATCHED-TEXT'])[:40]!r}"
        )
    print("--- EXOTIC NAMES (end) ---")
    for i in range(100, len(injected)):
        r = injected.iloc[i]
        print(f"  [{i}] {r['SOURCE']!r} D={r['DIALOGUE']!r}")
    print("--- ALL INJECTED-ERROR ---")
    for i in sorted(notes):
        print(f"  [{i}] {notes[i]}")


if __name__ == "__main__":
    main()
